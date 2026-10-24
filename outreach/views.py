import csv
from io import BytesIO

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.contrib.auth.views import redirect_to_login
from django.contrib.auth import SESSION_KEY
from django.contrib.sessions.models import Session
from django.core.paginator import Paginator
from django.core.exceptions import PermissionDenied, ValidationError
from django.db import IntegrityError, connection, transaction
from django.db.models.deletion import ProtectedError
from django.db.models import Case, CharField, Count, DateField, F, Q, Value, When
from django.db.models.functions import Coalesce, Lower
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.utils.http import url_has_allowed_host_and_scheme
from django.views.decorators.cache import never_cache
from django.views.decorators.http import require_http_methods, require_POST

from .bulk_import import (
    BulkImportFileError,
    import_prospects_from_csv,
    rollback_prospect_import,
)
from .forms import (
    BulkProspectImportForm,
    DashboardFilterForm,
    OutreachForm,
    ProspectForm,
    UserProvisionForm,
)
from .models import Outreach, Profile, Prospect, ProspectImportBatch
from .permissions import (
    active_role,
    assigned_roles,
    can_view_all_outreach,
    can_view_user_access,
    is_account_executive,
    is_ai_user,
    is_manager,
    is_submission_user,
    is_system_admin,
    role_membership_query,
)
from .workflows import sync_prospect_from_outreach
from .legacy_handoffs import ensure_handoff


REQUIRED_HEALTH_TABLES = {
    "auth_user",
    "django_migrations",
    "django_session",
    "django_site",
    "outreach_prospect",
    "outreach_escalation",
    "outreach_opportunity",
    "outreach_crmactivity",
    "outreach_profilerole",
    "outreach_foundercaeassignment",
    "socialaccount_socialaccount",
}


def login_landing(request):
    if request.user.is_authenticated:
        return redirect("dashboard")
    return render(request, "account/login.html")


@login_required
@never_cache
@require_http_methods(["GET", "POST"])
def choose_role(request):
    from .middleware import ACTIVE_ROLE_SESSION_KEY

    roles = assigned_roles(request.user)
    if len(roles) == 1:
        return redirect("dashboard")
    next_url = request.POST.get("next", "") if request.method == "POST" else request.GET.get("next", "")
    if not url_has_allowed_host_and_scheme(next_url, allowed_hosts={request.get_host()}, require_https=request.is_secure()):
        next_url = ""
    if next_url.split("?", 1)[0] in {request.path, "/login/"}:
        next_url = ""
    error = ""
    if request.method == "POST":
        selected = request.POST.get("role")
        if selected in roles:
            request.session.cycle_key()
            request.session[ACTIVE_ROLE_SESSION_KEY] = selected
            return redirect(next_url or "dashboard")
        error = "Choose one of the roles assigned to your account."
    return render(request, "account/choose_role.html", {
        "role_choices": request.user.profile.role_choices,
        "selected_role": active_role(request.user),
        "next_url": next_url,
        "role_error": error,
    }, status=400 if error else 200)


def health(request):
    try:
        with connection.cursor() as cursor:
            cursor.execute("SELECT 1")
            cursor.fetchone()
        available_tables = set(connection.introspection.table_names())
        if not REQUIRED_HEALTH_TABLES.issubset(available_tables):
            raise RuntimeError("Required database tables are unavailable.")
        return JsonResponse({"status": "ok"})
    except Exception:
        return JsonResponse({"status": "unavailable"}, status=503)


def scoped_prospects(user):
    queryset = Prospect.objects.select_related("owner", "owner__profile")
    if can_view_all_outreach(user):
        return queryset
    role = active_role(user)
    if role not in Prospect.Workstream.values:
        return queryset.none()
    return queryset.filter(workstream=role).filter(Q(owner=user) | Q(owner__isnull=True))


def get_scoped_prospect(user, pk):
    return get_object_or_404(scoped_prospects(user), pk=pk)


def scoped_company_rows(user):
    companies = {}
    for prospect in scoped_prospects(user).order_by("company_id", "pk"):
        companies.setdefault(
            prospect.company_id,
            (
                str(prospect.company_id),
                prospect.company_name,
                prospect.location,
                prospect.website,
            ),
        )
    return list(companies.values())


@login_required
@require_http_methods(["GET"])
def company_export(request, file_format):
    rows = scoped_company_rows(request.user)
    filename = f"flexgcc-companies-{timezone.localdate():%Y%m%d}"
    headers = ["ID", "Name", "Location", "URL"]

    if file_format == "csv":
        response = HttpResponse(content_type="text/csv; charset=utf-8")
        response["Content-Disposition"] = f'attachment; filename="{filename}.csv"'
        response["Cache-Control"] = "private, no-store"
        response.write("\ufeff")
        writer = csv.writer(response)
        writer.writerow(headers)
        writer.writerows(rows)
        return response

    if file_format == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Companies"
        worksheet.append(headers)
        for row in rows:
            worksheet.append(row)

        header_fill = PatternFill("solid", fgColor="087A55")
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(vertical="center")
        worksheet.row_dimensions[1].height = 24
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        worksheet.sheet_view.showGridLines = False
        worksheet.column_dimensions["A"].width = 39
        worksheet.column_dimensions["B"].width = 34
        worksheet.column_dimensions["C"].width = 28
        worksheet.column_dimensions["D"].width = 55
        for cell in worksheet["D"][1:]:
            if cell.value:
                cell.hyperlink = cell.value
                cell.style = "Hyperlink"

        output = BytesIO()
        workbook.save(output)
        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = f'attachment; filename="{filename}.xlsx"'
        response["Cache-Control"] = "private, no-store"
        return response

    return JsonResponse({"error": "Choose csv or xlsx."}, status=404)


@login_required
def dashboard(request, history=False):
    if is_account_executive(request.user) or (is_submission_user(request.user) and not history):
        return redirect("crm_lead_list")
    today = timezone.localdate()
    base_queryset = scoped_prospects(request.user)
    base_queryset = base_queryset.annotate(
        salesperson_id=Coalesce("owner_id", "created_by_id"),
        salesperson_name=Coalesce("owner__first_name", "created_by__first_name"),
        salesperson_email=Coalesce("owner__email", "created_by__email"),
        **{
            "next_owner_" + suffix: Case(
                When(escalation__opportunity__isnull=False, then=Coalesce("escalation__opportunity__next_action_owner" + field, "escalation__assigned_to" + field)),
                default=Coalesce("escalation__next_action_owner" + field, "escalation__assigned_to" + field, "next_action_owner" + field, "owner" + field),
            ) for suffix, field in [("id", "_id"), ("name", "__first_name"), ("email", "__email")]
        },
        effective_action=Case(
            When(Q(escalation__status="dropped") | Q(escalation__approval_status="pending") | Q(escalation__opportunity__stage__in=["won", "lost"]), then=Value("")),
            When(escalation__opportunity__isnull=False, then=F("escalation__opportunity__next_action")),
            When(escalation__isnull=False, then=F("escalation__next_action")),
            default=F("next_action"), output_field=CharField(),
        ),
        effective_action_date=Case(
            When(Q(escalation__status="dropped") | Q(escalation__approval_status="pending") | Q(escalation__opportunity__stage__in=["won", "lost"]), then=Value(None)),
            When(escalation__opportunity__isnull=False, then=F("escalation__opportunity__next_action_date")),
            When(escalation__isnull=False, then=F("escalation__next_action_date")),
            default=F("next_action_date"), output_field=DateField(),
        ),
    )
    summary = base_queryset.aggregate(
        total=Count("id"),
        unassigned=Count("id", filter=Q(owner__isnull=True)),
        actions_today=Count("id", filter=Q(effective_action_date=today)),
        overdue=Count("id", filter=Q(effective_action_date__lt=today)),
        meetings_scheduled=Count("id", filter=Q(status=Prospect.Status.MEETING_SCHEDULED)),
    )
    status_counts = {row["status"]: row["total"] for row in base_queryset.values("status").annotate(total=Count("id"))}
    stage_counts = {row["stage"]: row["total"] for row in base_queryset.values("stage").annotate(total=Count("id"))}
    workstream_counts = {
        row["workstream"]: row["total"] for row in base_queryset.values("workstream").annotate(total=Count("id"))
    }

    filter_form = DashboardFilterForm(request.GET or None, user=request.user)
    prospects = base_queryset.annotate(outreach_count=Count("outreaches"))
    if filter_form.is_valid():
        data = filter_form.cleaned_data
        if data.get("search"):
            search = data["search"].strip()
            prospects = prospects.filter(
                Q(company_name__icontains=search)
                | Q(contact_name__icontains=search)
                | Q(contact_email__icontains=search)
                | Q(location__icontains=search)
                | Q(consulting_focus__icontains=search)
            )
        if data.get("status"):
            prospects = prospects.filter(status=data["status"])
        if data.get("stage"):
            prospects = prospects.filter(stage=data["stage"])
        if data.get("claim_status") == "claimed":
            prospects = prospects.filter(owner__isnull=False)
        elif data.get("claim_status") == "unclaimed":
            prospects = prospects.filter(owner__isnull=True)
        if data.get("workstream"):
            prospects = prospects.filter(workstream=data["workstream"])
        if data.get("outreach_count") != "" and data.get("outreach_count") is not None:
            prospects = prospects.filter(outreach_count=int(data["outreach_count"]))
        if data.get("due") == "today":
            prospects = prospects.filter(effective_action_date=today)
        elif data.get("due") == "overdue":
            prospects = prospects.filter(effective_action_date__lt=today)
        elif data.get("due") == "upcoming":
            prospects = prospects.filter(effective_action_date__gt=today)
        elif data.get("due") == "none":
            prospects = prospects.filter(effective_action_date__isnull=True)
        if data.get("owner"):
            prospects = prospects.filter(salesperson_id=data["owner"].pk)
        if data.get("next_owner"):
            prospects = prospects.filter(next_owner_id=data["next_owner"].pk)

    prospects = prospects.order_by(F("effective_action_date").asc(nulls_last=True), "company_name")
    if filter_form.is_valid() and filter_form.cleaned_data.get("sort"):
        sort = filter_form.cleaned_data["sort"]
        name = Lower(sort.lstrip("-") + "_name")
        email = Lower(sort.lstrip("-") + "_email")
        prospects = prospects.order_by(name.desc(nulls_last=True) if sort.startswith("-") else name.asc(nulls_last=True), email.desc() if sort.startswith("-") else email.asc(), "company_name", "pk")
    paginator = Paginator(prospects, 20)
    page_obj = paginator.get_page(request.GET.get("page"))
    query_params = request.GET.copy()
    query_params.pop("page", None)

    context = {
        "summary": summary,
        "status_counts": status_counts,
        "status_summary": [
            (value, label, status_counts.get(value, 0)) for value, label in Prospect.Status.choices
        ],
        "stage_summary": [
            (value, label, stage_counts.get(value, 0)) for value, label in Prospect.Stage.choices
        ],
        "workstream_summary": [
            (value, label, workstream_counts.get(value, 0)) for value, label in Prospect.Workstream.choices
        ],
        "status_choices": Prospect.Status.choices,
        "filter_form": filter_form,
        "page_obj": page_obj,
        "query_string": query_params.urlencode(),
        "today": today,
        "operator_count": (
            Profile.objects.filter(role_membership_query(Prospect.Workstream.values, prefix=""), user__is_active=True).distinct().count()
            if can_view_all_outreach(request.user)
            else None
        ),
    }
    return render(request, "outreach/dashboard.html", context)


@login_required
@require_http_methods(["GET", "POST"])
def prospect_create(request):
    if is_ai_user(request.user):
        raise PermissionDenied("AI user access is read-only.")
    # Keep old bookmarks, but never offer a second new-record workflow.
    if request.method == "POST":
        from .crm_views import submit
        return submit(request)
    return redirect("crm_submit")


def _get_scoped_import_batch(user, batch_id):
    batches = ProspectImportBatch.objects.select_related(
        "uploaded_by",
        "owner",
        "rolled_back_by",
    )
    if not can_view_all_outreach(user):
        batches = batches.filter(uploaded_by=user, target_workstream=active_role(user))
    return get_object_or_404(batches, pk=batch_id)


@login_required
@require_http_methods(["GET", "POST"])
def prospect_bulk_import(request):
    if is_ai_user(request.user):
        raise PermissionDenied("AI user access is read-only.")
    form = BulkProspectImportForm(
        request.POST or None,
        request.FILES or None,
        user=request.user,
    )
    if request.method == "POST" and form.is_valid():
        try:
            batch = import_prospects_from_csv(
                uploaded_file=form.cleaned_data["csv_file"],
                uploaded_by=request.user,
                workstream=form.cleaned_data["workstream"],
                owner=form.cleaned_data["owner"],
            )
        except BulkImportFileError as error:
            form.add_error("csv_file", str(error))
        else:
            messages.success(
                request,
                f"CSV ingestion completed: {batch.created_count} added and "
                f"{batch.skipped_count} ignored.",
            )
            return redirect("prospect_bulk_import_result", batch_id=batch.pk)
    recent_batches = ProspectImportBatch.objects.select_related("owner", "uploaded_by")
    if not is_manager(request.user):
        recent_batches = recent_batches.filter(uploaded_by=request.user, target_workstream=active_role(request.user))
    response = render(
        request,
        "outreach/prospect_bulk_import.html",
        {
            "form": form,
            "recent_batches": recent_batches[:10],
        },
    )
    response["Cache-Control"] = "private, no-store"
    return response


@login_required
@require_http_methods(["GET"])
def prospect_bulk_import_result(request, batch_id):
    batch = _get_scoped_import_batch(request.user, batch_id)
    current_by_id = {
        prospect.pk: prospect
        for prospect in batch.prospects.select_related("owner").order_by("company_name")
    }
    created_rows = []
    for snapshot in batch.created_rows:
        row = dict(snapshot)
        row["prospect"] = current_by_id.get(snapshot["prospect_id"])
        created_rows.append(row)
    response = render(
        request,
        "outreach/prospect_bulk_import_result.html",
        {
            "batch": batch,
            "created_rows": created_rows,
            "remaining_count": len(current_by_id),
            "can_rollback": not is_ai_user(request.user) and bool(current_by_id)
            and batch.status != ProspectImportBatch.Status.ROLLED_BACK,
        },
    )
    response["Cache-Control"] = "private, no-store"
    return response


@login_required
@require_POST
def prospect_bulk_import_rollback(request, batch_id):
    batch = _get_scoped_import_batch(request.user, batch_id)
    if batch.status == ProspectImportBatch.Status.ROLLED_BACK:
        messages.info(request, "This import batch has already been rolled back.")
        return redirect("prospect_bulk_import_result", batch_id=batch.pk)

    batch, removed_count, protected_rows = rollback_prospect_import(
        batch_id=batch.pk,
        rolled_back_by=request.user,
    )
    if protected_rows:
        messages.warning(
            request,
            f"Rollback removed {removed_count} prospect(s). "
            f"{len(protected_rows)} modified or active prospect(s) were protected.",
        )
    else:
        messages.success(
            request,
            f"Rollback completed. {removed_count} imported prospect(s) were removed.",
        )
    return redirect("prospect_bulk_import_result", batch_id=batch.pk)


@login_required
@require_http_methods(["GET", "POST"])
@transaction.atomic
def prospect_update(request, pk):
    if is_ai_user(request.user):
        raise PermissionDenied("AI user access is read-only.")
    prospect = get_object_or_404(
        scoped_prospects(request.user).select_related(None).select_for_update(), pk=pk
    )
    if not is_manager(request.user) and prospect.owner_id != request.user.id:
        messages.error(request, "Claim this prospect before editing its research or follow-up details.")
        return redirect(prospect)
    form = ProspectForm(request.POST or None, instance=prospect, user=request.user)
    if request.method == "POST" and form.is_valid():
        try:
            prospect = form.save()
        except ValidationError as error:
            form.add_error(None, "; ".join(error.messages))
        else:
            messages.success(request, f"{prospect.company_name} was updated.")
            return redirect(prospect)
    return render(
        request,
        "outreach/prospect_form.html",
        {"form": form, "prospect": prospect, "title": f"Edit {prospect.company_name}"},
    )


@login_required
def prospect_detail(request, pk):
    prospect = get_scoped_prospect(request.user, pk)
    outreaches = prospect.outreaches.select_related("recorded_by").all()
    can_edit_prospect = not is_ai_user(request.user) and (
        is_manager(request.user) or prospect.owner_id == request.user.id
    )
    return render(
        request,
        "outreach/prospect_detail.html",
        {
            "prospect": prospect,
            "outreaches": outreaches,
            "outreach_form": OutreachForm(prospect=prospect),
            "can_add_outreach": can_edit_prospect and outreaches.count() < 5,
            "can_edit_prospect": can_edit_prospect,
            "can_claim_prospect": (
                not is_manager(request.user)
                and prospect.owner_id is None
                and prospect.workstream == active_role(request.user)
            ),
            "today": timezone.localdate(),
        },
    )


@login_required
@require_POST
def outreach_add(request, pk):
    with transaction.atomic():
        # scoped_prospects() joins nullable owner for display. PostgreSQL rejects
        # FOR UPDATE on that outer join, so lock only the prospect table here.
        prospect = get_object_or_404(
            scoped_prospects(request.user).select_related(None).select_for_update(),
            pk=pk,
        )
        if prospect.owner_id != request.user.id and not is_manager(request.user):
            messages.error(request, "Claim this prospect before recording outreach.")
            return redirect(prospect)
        existing_numbers = list(prospect.outreaches.order_by("sequence_number").values_list("sequence_number", flat=True))
        if len(existing_numbers) >= 5:
            messages.error(request, "This prospect already has the maximum of five outreach records.")
            return redirect(prospect)
        form = OutreachForm(request.POST, prospect=prospect)
        if form.is_valid():
            outreach = form.save(commit=False)
            outreach.prospect = prospect
            outreach.recorded_by = request.user
            outreach.sequence_number = next(number for number in range(1, 6) if number not in existing_numbers)
            try:
                with transaction.atomic():
                    outreach.save()
                    sync_prospect_from_outreach(prospect, outreach, request.user)
            except IntegrityError:
                messages.error(request, "Another outreach was added at the same time. Please try again.")
            except ValidationError as error:
                form.add_error(None, "; ".join(error.messages))
            else:
                messages.success(request, f"Outreach {outreach.sequence_number} was recorded.")
                return redirect(prospect)
    outreaches = prospect.outreaches.select_related("recorded_by").all()
    return render(
        request,
        "outreach/prospect_detail.html",
        {
            "prospect": prospect,
            "outreaches": outreaches,
            "outreach_form": form,
            "can_add_outreach": outreaches.count() < 5,
            "can_edit_prospect": is_manager(request.user) or prospect.owner_id == request.user.id,
            "can_claim_prospect": False,
            "today": timezone.localdate(),
        },
        status=400,
    )


@login_required
@require_POST
def prospect_claim(request, pk):
    if is_manager(request.user):
        messages.error(request, "Managers can assign prospects from the edit screen.")
        return redirect("dashboard")

    with transaction.atomic():
        prospect = get_object_or_404(Prospect.objects.select_for_update(), pk=pk)
        if prospect.workstream != active_role(request.user):
            raise PermissionDenied("This prospect belongs to another outreach workstream.")
        if prospect.owner_id is None:
            prospect.owner = request.user
            prospect.save(update_fields=["owner", "updated_at"])
            messages.success(request, f"{prospect.company_name} is now assigned to you.")
        elif prospect.owner_id == request.user.id:
            messages.info(request, f"{prospect.company_name} is already assigned to you.")
        else:
            messages.error(request, "Another user has already claimed this prospect.")
            return redirect("dashboard")
    return redirect(prospect)


@login_required
@require_http_methods(["GET", "POST"])
@transaction.atomic
def outreach_update(request, pk):
    if is_ai_user(request.user):
        raise PermissionDenied("AI user access is read-only.")
    outreach = get_object_or_404(Outreach.objects.select_related("prospect"), pk=pk)
    prospect = get_scoped_prospect(request.user, outreach.prospect_id)
    prospect = Prospect.objects.select_for_update().get(pk=prospect.pk)
    outreach = Outreach.objects.select_for_update().get(pk=pk)
    if prospect.owner_id != request.user.id and not is_manager(request.user):
        raise PermissionDenied("Claim this prospect before editing outreach history.")
    form = OutreachForm(request.POST or None, instance=outreach, prospect=prospect)
    if request.method == "POST" and form.is_valid():
        try:
            with transaction.atomic():
                outreach = form.save()
                sync_prospect_from_outreach(prospect, outreach, request.user)
        except ValidationError as error:
            form.add_error(None, "; ".join(error.messages))
        else:
            messages.success(request, f"Outreach {outreach.sequence_number} was updated.")
            return redirect(prospect)
    return render(
        request,
        "outreach/outreach_form.html",
        {"form": form, "outreach": outreach, "prospect": prospect},
    )


def home(request):
    if request.user.is_authenticated:
        return redirect("dashboard")
    return redirect_to_login("/dashboard/", login_url="login")


def require_system_admin(user):
    if not is_system_admin(user):
        raise PermissionDenied("System administrator access is required.")


def invalidate_user_sessions(user):
    for session in Session.objects.filter(expire_date__gte=timezone.now()).iterator():
        if session.get_decoded().get(SESSION_KEY) == str(user.pk):
            session.delete()


@login_required
def user_list(request):
    if not can_view_user_access(request.user):
        raise PermissionDenied("System administration visibility is required.")
    users = (
        Profile.objects.select_related("user")
        .prefetch_related("additional_roles")
        .order_by("-user__is_active", "role", "user__first_name", "user__email")
    )
    counts = {
        "active": users.filter(user__is_active=True).count(),
    }
    for key, role in {"interns": Profile.Role.INTERN, "inside_sales": Profile.Role.INSIDE_SALES,
                      "linkedin_outreach": Profile.Role.LINKEDIN_OUTREACH, "executives": Profile.Role.ACCOUNT_EXECUTIVE,
                      "managers": Profile.Role.MANAGER, "admins": Profile.Role.SYSTEM_ADMIN,
                      "ai_users": Profile.Role.AI_USER}.items():
        counts[key] = users.filter(role_membership_query([role], prefix=""), user__is_active=True).distinct().count()
    return render(request, "outreach/user_list.html", {"profiles": users, "counts": counts})


@login_required
@require_http_methods(["GET", "POST"])
def user_create(request):
    require_system_admin(request.user)
    form = UserProvisionForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        user = form.save()
        messages.success(request, f"{user.get_full_name()} was added and can now sign in with {user.email}.")
        return redirect("user_list")
    return render(request, "outreach/user_form.html", {"form": form, "title": "Add user"})


@login_required
@require_http_methods(["GET", "POST"])
def user_update(request, pk):
    require_system_admin(request.user)
    target = get_object_or_404(Profile.objects.select_related("user"), user_id=pk).user
    form = UserProvisionForm(request.POST or None, instance=target)
    if request.method == "POST" and form.is_valid():
        new_roles = form.cleaned_data["roles"]
        if target == request.user and Profile.Role.SYSTEM_ADMIN not in new_roles:
            form.add_error("roles", "You cannot remove your own system administrator access.")
        elif (
            target.profile.has_role(Profile.Role.SYSTEM_ADMIN)
            and Profile.Role.SYSTEM_ADMIN not in new_roles
            and not Profile.objects.filter(role_membership_query([Profile.Role.SYSTEM_ADMIN], prefix=""), user__is_active=True).exclude(user=target).exists()
        ):
            form.add_error("roles", "Add another active system administrator before changing the last one.")
        else:
            user = form.save()
            messages.success(request, f"{user.get_full_name()} was updated.")
            return redirect("user_list")
    return render(request, "outreach/user_form.html", {"form": form, "target_user": target, "title": "Edit user"})


@login_required
@require_http_methods(["GET", "POST"])
def user_delete(request, pk):
    require_system_admin(request.user)
    target = get_object_or_404(Profile.objects.select_related("user"), user_id=pk).user
    if target == request.user:
        messages.error(request, "You cannot delete your own account.")
        return redirect("user_list")
    if (
        target.profile.has_role(Profile.Role.SYSTEM_ADMIN)
        and target.is_active
        and not Profile.objects.filter(role_membership_query([Profile.Role.SYSTEM_ADMIN], prefix=""), user__is_active=True).exclude(user=target).exists()
    ):
        messages.error(request, "Add another active system administrator before deleting the last one.")
        return redirect("user_list")
    if request.method == "POST":
        name = target.get_full_name() or target.email
        try:
            target.delete()
        except ProtectedError:
            target.is_active = False
            target.save(update_fields=["is_active"])
            invalidate_user_sessions(target)
            messages.warning(request, f"{name}'s access was removed. Their account record was retained to preserve sales history.")
        else:
            messages.success(request, f"{name} was deleted.")
        return redirect("user_list")
    return render(request, "outreach/user_confirm_delete.html", {"target_user": target})


@login_required
@require_POST
def user_restore(request, pk):
    require_system_admin(request.user)
    target = get_object_or_404(Profile.objects.select_related("user"), user_id=pk).user
    target.is_active = True
    target.save(update_fields=["is_active"])
    messages.success(request, f"{target.get_full_name() or target.email} can sign in again.")
    return redirect("user_list")
