import csv
import hashlib
import json
import uuid
from datetime import date, timedelta
from html import escape
from io import BytesIO, StringIO

import jwt
from django.conf import settings
from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.db import connection
from django.test import TestCase
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from outreach.api_tokens import (
    API_ACCESS_TOKEN_PREFIX,
    ApiTokenError,
    authenticate_access_token,
    issue_token_pair,
)
from outreach.models import (
    ApiAccessToken,
    ApiRefreshToken,
    CompanyUpdateAudit,
    Outreach,
    OutreachUpdateAudit,
    Profile,
    Prospect,
    ProspectUpdateAudit,
)


class CompanyApiTestMixin:
    def make_user(self, email, role=Profile.Role.INTERN, name="Test User"):
        first_name, _, last_name = name.partition(" ")
        user = get_user_model().objects.create_user(
            username=email,
            email=email,
            first_name=first_name,
            last_name=last_name,
        )
        user.profile.role = role
        user.profile.save(update_fields=["role"])
        return user

    def make_prospect(self, owner, company_name, website, **overrides):
        values = {
            "owner": owner,
            "workstream": owner.profile.role,
            "stage": Prospect.Stage.ELIGIBLE,
            "company_name": company_name,
            "website": website,
            "short_description": "Boutique advisory firm.",
            "location": "Chicago, IL",
            "consulting_focus": "Operations consulting",
            "client_segment": "Mid-market companies",
            "eligibility_evidence": "Published mid-market client work.",
            "contact_name": "Asha Rao",
            "contact_title": "Managing Partner",
            "contact_email": "asha@example.com",
            "created_by": overrides.pop("created_by", owner),
        }
        values.update(overrides)
        return Prospect.objects.create(**values)

    def bearer(self, user):
        return f"Bearer {issue_token_pair(user)['access_token']}"


class CompanyIdentityTests(CompanyApiTestMixin, TestCase):
    def test_same_domain_shares_one_system_generated_company_id_across_workstreams(self):
        intern = self.make_user("intern@example.com")
        inside_sales = self.make_user("inside@example.com", Profile.Role.INSIDE_SALES)
        first = self.make_prospect(intern, "Shared Advisory", "https://shared.example.com")
        second = self.make_prospect(
            inside_sales,
            "Shared Advisory",
            "https://www.shared.example.com/about",
        )

        self.assertEqual(first.company_id, second.company_id)
        self.assertIsInstance(first.company_id, uuid.UUID)

    def test_company_id_cannot_be_changed_after_save(self):
        intern = self.make_user("intern@example.com")
        prospect = self.make_prospect(intern, "Stable ID Advisory", "https://stable.example.com")
        prospect.company_id = uuid.uuid4()

        with self.assertRaises(ValidationError) as error:
            prospect.save()

        self.assertIn("cannot be changed", error.exception.message_dict["company"][0])


class CompanyExportTests(CompanyApiTestMixin, TestCase):
    def setUp(self):
        self.intern = self.make_user("intern@example.com", name="Isha Intern")
        self.inside = self.make_user("inside@example.com", Profile.Role.INSIDE_SALES, "Ivan Inside")
        self.manager = self.make_user("manager@example.com", Profile.Role.MANAGER, "Meera Manager")
        self.first = self.make_prospect(
            self.intern,
            "Shared Advisory",
            "https://shared.example.com",
            location="Miami, FL",
        )
        self.make_prospect(
            self.inside,
            "Shared Advisory",
            "https://www.shared.example.com/about",
            location="Miami, FL",
        )

    def test_csv_export_has_exact_columns_and_deduplicates_company_ids(self):
        self.client.force_login(self.manager)
        response = self.client.get(reverse("company_export", args=["csv"]))

        self.assertEqual(response.status_code, 200)
        rows = list(csv.reader(StringIO(response.content.decode("utf-8-sig"))))
        self.assertEqual(rows[0], ["ID", "Name", "Location", "URL"])
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1][0], str(self.first.company_id))
        self.assertEqual(rows[1][1:], ["Shared Advisory", "Miami, FL", "https://shared.example.com"])

    def test_xlsx_export_is_scoped_and_formatted_for_the_logged_in_user(self):
        self.client.force_login(self.intern)
        response = self.client.get(reverse("company_export", args=["xlsx"]))

        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content), read_only=False)
        worksheet = workbook["Companies"]
        self.assertEqual([cell.value for cell in worksheet[1]], ["ID", "Name", "Location", "URL"])
        self.assertEqual(worksheet.max_row, 2)
        self.assertEqual(worksheet["A2"].value, str(self.first.company_id))
        self.assertEqual(worksheet.freeze_panes, "A2")
        self.assertEqual(worksheet.auto_filter.ref, "A1:D2")


class TokenLifecycleTests(CompanyApiTestMixin, TestCase):
    def setUp(self):
        self.user = self.make_user("intern@example.com", name="Isha Intern")

    def test_logged_in_user_generates_persistent_hashed_access_token(self):
        self.client.force_login(self.user)
        response = self.client.post(reverse("api_access"))

        self.assertEqual(response.status_code, 200)
        token_data = response.context["token_pair"]
        raw_token = token_data["access_token"]
        self.assertTrue(raw_token.startswith(API_ACCESS_TOKEN_PREFIX))
        self.assertIsNone(token_data["expires_in"])
        self.assertNotIn("refresh_token", token_data)

        stored = ApiAccessToken.objects.get(user=self.user)
        self.assertNotEqual(stored.token_hash, raw_token)
        self.assertEqual(
            stored.token_hash,
            hashlib.sha256(raw_token.encode("utf-8")).hexdigest(),
        )
        self.assertEqual(stored.token_prefix, raw_token[:16])
        self.assertContains(response, "No automatic expiry")

    def test_persistent_token_does_not_expire_based_on_age(self):
        raw_token = issue_token_pair(self.user)["access_token"]
        token = ApiAccessToken.objects.get(user=self.user)
        ApiAccessToken.objects.filter(pk=token.pk).update(
            created_at=timezone.now() - timedelta(days=3650)
        )

        authenticated_user = authenticate_access_token(raw_token)

        self.assertEqual(authenticated_user, self.user)
        token.refresh_from_db()
        self.assertIsNotNone(token.last_used_at)

    def test_user_can_revoke_own_token_and_it_stops_working_immediately(self):
        prospect = self.make_prospect(
            self.user,
            "Persistent Token Advisory",
            "https://persistent-token.example.com",
        )
        raw_token = issue_token_pair(self.user)["access_token"]
        token = ApiAccessToken.objects.get(user=self.user)
        company_url = reverse("api_company_detail", args=[prospect.company_id])
        self.assertEqual(
            self.client.get(
                company_url,
                HTTP_AUTHORIZATION=f"Bearer {raw_token}",
            ).status_code,
            200,
        )

        self.client.force_login(self.user)
        revoke_response = self.client.post(
            reverse("api_access_token_revoke", args=[token.pk])
        )

        self.assertRedirects(revoke_response, reverse("api_access"))
        token.refresh_from_db()
        self.assertIsNotNone(token.revoked_at)
        rejected = self.client.get(
            company_url,
            HTTP_AUTHORIZATION=f"Bearer {raw_token}",
        )
        self.assertEqual(rejected.status_code, 401)
        self.assertIn("revoked", rejected.json()["error"])

    def test_user_cannot_revoke_another_users_token(self):
        other_user = self.make_user("other@example.com", name="Omar Intern")
        issue_token_pair(other_user)
        other_token = ApiAccessToken.objects.get(user=other_user)
        self.client.force_login(self.user)

        response = self.client.post(
            reverse("api_access_token_revoke", args=[other_token.pk])
        )

        self.assertEqual(response.status_code, 404)
        other_token.refresh_from_db()
        self.assertIsNone(other_token.revoked_at)

    def test_token_is_rejected_when_its_user_is_inactive(self):
        raw_token = issue_token_pair(self.user)["access_token"]
        self.user.is_active = False
        self.user.save(update_fields=["is_active"])

        with self.assertRaisesMessage(ApiTokenError, "unavailable or inactive"):
            authenticate_access_token(raw_token)

    def test_old_refresh_token_is_exchanged_once_for_persistent_token(self):
        raw_refresh_token = "old-refresh-token-value"
        ApiRefreshToken.objects.create(
            user=self.user,
            token_hash=hashlib.sha256(raw_refresh_token.encode("utf-8")).hexdigest(),
            expires_at=timezone.now() + timedelta(days=1),
        )
        response = self.client.post(
            reverse("api_token_refresh"),
            data=json.dumps({"refresh_token": raw_refresh_token}),
            content_type="application/json",
        )

        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()["access_token"].startswith(API_ACCESS_TOKEN_PREFIX))
        self.assertIsNone(response.json()["expires_in"])
        self.assertNotIn("refresh_token", response.json())
        self.assertEqual(ApiAccessToken.objects.filter(user=self.user).count(), 1)

        reused = self.client.post(
            reverse("api_token_refresh"),
            data=json.dumps({"refresh_token": raw_refresh_token}),
            content_type="application/json",
        )
        self.assertEqual(reused.status_code, 401)

    def test_unexpired_legacy_jwt_remains_compatible(self):
        now = timezone.now()
        legacy_token = jwt.encode(
            {
                "iss": settings.API_TOKEN_ISSUER,
                "aud": settings.API_TOKEN_AUDIENCE,
                "sub": str(self.user.pk),
                "email": self.user.email,
                "type": "access",
                "jti": uuid.uuid4().hex,
                "iat": int(now.timestamp()),
                "exp": int((now + timedelta(minutes=15)).timestamp()),
            },
            settings.API_TOKEN_SIGNING_KEY,
            algorithm="HS256",
        )

        self.assertEqual(authenticate_access_token(legacy_token), self.user)


class CompanyApiTests(CompanyApiTestMixin, TestCase):
    def setUp(self):
        self.intern = self.make_user("intern@example.com", name="Isha Intern")
        self.other_intern = self.make_user("other@example.com", name="Omar Intern")
        self.inside = self.make_user("inside@example.com", Profile.Role.INSIDE_SALES, "Ivan Inside")
        self.manager = self.make_user("manager@example.com", Profile.Role.MANAGER, "Meera Manager")
        self.first = self.make_prospect(
            self.intern,
            "Shared Advisory",
            "https://shared.example.com",
        )
        self.second = self.make_prospect(
            self.inside,
            "Shared Advisory",
            "https://www.shared.example.com/about",
        )

    def patch_company(self, user, payload):
        return self.client.generic(
            "PATCH",
            reverse("api_company_detail", args=[self.first.company_id]),
            data=json.dumps(payload),
            content_type="application/json",
            HTTP_AUTHORIZATION=self.bearer(user),
        )

    def test_manager_update_propagates_to_all_workstreams_and_is_audited(self):
        response = self.patch_company(
            self.manager,
            {
                "short_description": "Updated canonical description.",
                "eligibility_evidence": "Verified boutique consultancy.",
                "linkedin_url": "https://www.linkedin.com/in/asha-rao",
                "email": "new@example.com",
                "phone": "+1 312 555 0100",
            },
        )

        self.assertEqual(response.status_code, 200, response.content)
        self.first.refresh_from_db()
        self.second.refresh_from_db()
        for prospect in (self.first, self.second):
            self.assertEqual(prospect.short_description, "Updated canonical description.")
            self.assertEqual(prospect.contact_email, "new@example.com")
            self.assertEqual(prospect.contact_phone, "+1 312 555 0100")
        audit = CompanyUpdateAudit.objects.get(company_id=self.first.company_id)
        self.assertEqual(audit.modified_by, self.manager)
        self.assertEqual(audit.changed_values["email"], "new@example.com")
        self.assertEqual(response.json()["modified_by"]["name"], "Meera Manager")

    def test_company_id_and_unsupported_fields_cannot_be_modified_by_api(self):
        original_id = self.first.company_id
        response = self.patch_company(self.manager, {"id": str(uuid.uuid4())})

        self.assertEqual(response.status_code, 400)
        self.first.refresh_from_db()
        self.assertEqual(self.first.company_id, original_id)
        self.assertEqual(CompanyUpdateAudit.objects.count(), 0)

    def test_access_token_is_required(self):
        response = self.client.generic(
            "PATCH",
            reverse("api_company_detail", args=[self.first.company_id]),
            data=json.dumps({"contact_name": "Updated"}),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 401)

    def test_frontline_user_must_claim_before_api_update(self):
        unassigned = Prospect.objects.create(
            owner=None,
            workstream=Prospect.Workstream.INTERN,
            stage=Prospect.Stage.RESEARCH,
            company_name="Unassigned Advisory",
            website="https://unassigned.example.com",
            created_by=self.manager,
        )
        url = reverse("api_company_detail", args=[unassigned.company_id])
        before_claim = self.client.generic(
            "PATCH",
            url,
            data=json.dumps({"contact_name": "New Contact"}),
            content_type="application/json",
            HTTP_AUTHORIZATION=self.bearer(self.other_intern),
        )
        self.assertEqual(before_claim.status_code, 403)

        claim = self.client.post(
            reverse("api_company_claim", args=[unassigned.company_id]),
            HTTP_AUTHORIZATION=self.bearer(self.other_intern),
        )
        self.assertEqual(claim.status_code, 200)
        updated = self.client.generic(
            "PATCH",
            url,
            data=json.dumps({"contact_name": "New Contact"}),
            content_type="application/json",
            HTTP_AUTHORIZATION=self.bearer(self.other_intern),
        )
        self.assertEqual(updated.status_code, 200, updated.content)

    def test_claim_conflict_returns_existing_users_name(self):
        claimed = Prospect.objects.create(
            owner=self.intern,
            workstream=Prospect.Workstream.INTERN,
            stage=Prospect.Stage.RESEARCH,
            company_name="Claimed Advisory",
            website="https://claimed.example.com",
            created_by=self.manager,
        )
        response = self.client.post(
            reverse("api_company_claim", args=[claimed.company_id]),
            HTTP_AUTHORIZATION=self.bearer(self.other_intern),
        )

        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.json()["claimed_by"], "Isha Intern")
        self.assertIn("already claimed by Isha Intern", response.json()["error"])

    def test_claim_row_lock_does_not_join_nullable_owner(self):
        unassigned = Prospect.objects.create(
            owner=None,
            workstream=Prospect.Workstream.INTERN,
            stage=Prospect.Stage.RESEARCH,
            company_name="ORENG Consulting",
            website="http://www.orengconsulting.com",
            location="Boston, MA",
            created_by=self.manager,
        )

        with CaptureQueriesContext(connection) as queries:
            response = self.client.post(
                reverse("api_company_claim", args=[unassigned.company_id]),
                HTTP_AUTHORIZATION=self.bearer(self.other_intern),
            )

        self.assertEqual(response.status_code, 200, response.content)
        claim_selects = [
            query["sql"]
            for query in queries.captured_queries
            if 'FROM "outreach_prospect"' in query["sql"]
            and '"outreach_prospect"."company_id"' in query["sql"]
            and query["sql"].lstrip().upper().startswith("SELECT")
        ]
        self.assertTrue(claim_selects)
        for claim_select in claim_selects:
            self.assertNotIn('JOIN "auth_user"', claim_select)


class ProspectWorkflowApiTests(CompanyApiTestMixin, TestCase):
    def setUp(self):
        self.linkedin_user = self.make_user(
            "linkedin@example.com",
            Profile.Role.LINKEDIN_OUTREACH,
            "Leena LinkedIn",
        )
        self.other_linkedin_user = self.make_user(
            "other-linkedin@example.com",
            Profile.Role.LINKEDIN_OUTREACH,
            "Omar LinkedIn",
        )
        self.intern = self.make_user("intern@example.com", name="Isha Intern")
        self.manager = self.make_user(
            "manager@example.com",
            Profile.Role.MANAGER,
            "Meera Manager",
        )
        self.prospect = self.make_prospect(
            self.linkedin_user,
            "LinkedIn Advisory",
            "https://linkedin-advisory.example.com",
            contact_linkedin_url="https://www.linkedin.com/in/asha-rao",
            founder_account=Prospect.FounderAccount.KANDARP_SONI,
            linkedin_connection_status=Prospect.LinkedInConnectionStatus.NOT_SENT,
            personalization_note="Chicago operations practice.",
        )
        self.url = reverse("api_prospect_detail", args=[self.prospect.pk])

    def patch_prospect(self, user, payload, *, prospect=None):
        prospect = prospect or self.prospect
        return self.client.generic(
            "PATCH",
            reverse("api_prospect_detail", args=[prospect.pk]),
            data=json.dumps(payload),
            content_type="application/json",
            HTTP_AUTHORIZATION=self.bearer(user),
        )

    def test_manager_updates_complete_founder_section_and_new_flags_with_audit(self):
        response = self.patch_prospect(
            self.manager,
            {
                "founder_account": Prospect.FounderAccount.SUNIT_KALA,
                "linkedin_connection_status": Prospect.LinkedInConnectionStatus.REQUEST_SENT,
                "personalization_note": "Florida healthcare advisory practice.",
                "founder_escalation_required": True,
                "founder_escalation_notes": "Founder should answer the pricing question.",
                "prospect_sent": True,
                "is_not_eligible": True,
            },
        )

        self.assertEqual(response.status_code, 200, response.content)
        self.prospect.refresh_from_db()
        self.assertEqual(self.prospect.founder_account, Prospect.FounderAccount.SUNIT_KALA)
        self.assertEqual(
            self.prospect.linkedin_connection_status,
            Prospect.LinkedInConnectionStatus.REQUEST_SENT,
        )
        self.assertTrue(self.prospect.founder_escalation_required)
        self.assertTrue(self.prospect.prospect_sent)
        self.assertTrue(self.prospect.is_not_eligible)
        self.assertEqual(response.json()["modified_by"]["name"], "Meera Manager")
        audit = ProspectUpdateAudit.objects.get(prospect=self.prospect)
        self.assertEqual(audit.modified_by, self.manager)
        self.assertEqual(
            audit.previous_values["founder_account"],
            Prospect.FounderAccount.KANDARP_SONI,
        )
        self.assertTrue(audit.changed_values["prospect_sent"])

    def test_owner_can_update_and_another_frontline_user_cannot(self):
        updated = self.patch_prospect(
            self.linkedin_user,
            {"prospect_sent": True, "is_not_eligible": False},
        )
        rejected = self.patch_prospect(
            self.other_linkedin_user,
            {"prospect_sent": False},
        )

        self.assertEqual(updated.status_code, 200, updated.content)
        self.assertEqual(rejected.status_code, 403)
        self.prospect.refresh_from_db()
        self.assertTrue(self.prospect.prospect_sent)

    def test_non_linkedin_prospect_accepts_flags_but_rejects_founder_fields(self):
        intern_prospect = self.make_prospect(
            self.intern,
            "Intern Advisory",
            "https://intern-advisory.example.com",
        )
        flags_response = self.patch_prospect(
            self.intern,
            {"prospect_sent": True, "is_not_eligible": True},
            prospect=intern_prospect,
        )
        founder_response = self.patch_prospect(
            self.intern,
            {"founder_account": Prospect.FounderAccount.KANDARP_SONI},
            prospect=intern_prospect,
        )

        self.assertEqual(flags_response.status_code, 200, flags_response.content)
        self.assertEqual(founder_response.status_code, 400)
        self.assertIn(
            "founder_account",
            founder_response.json()["field_errors"],
        )
        intern_prospect.refresh_from_db()
        self.assertTrue(intern_prospect.prospect_sent)
        self.assertTrue(intern_prospect.is_not_eligible)
        self.assertEqual(intern_prospect.founder_account, "")

    def test_invalid_choice_and_string_boolean_return_field_errors(self):
        response = self.patch_prospect(
            self.manager,
            {
                "founder_account": "unknown_founder",
                "prospect_sent": "true",
            },
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn("founder_account", response.json()["field_errors"])
        self.assertIn("prospect_sent", response.json()["field_errors"])
        self.assertEqual(ProspectUpdateAudit.objects.count(), 0)

    def test_get_returns_workflow_fields_and_company_lists_prospect_ids(self):
        prospect_response = self.client.get(
            self.url,
            HTTP_AUTHORIZATION=self.bearer(self.manager),
        )
        company_response = self.client.get(
            reverse("api_company_detail", args=[self.prospect.company_id]),
            HTTP_AUTHORIZATION=self.bearer(self.manager),
        )

        self.assertEqual(prospect_response.status_code, 200)
        self.assertEqual(
            prospect_response.json()["prospect"]["founder_account"],
            Prospect.FounderAccount.KANDARP_SONI,
        )
        self.assertFalse(prospect_response.json()["prospect"]["prospect_sent"])
        self.assertFalse(prospect_response.json()["prospect"]["is_not_eligible"])
        self.assertIn(
            {
                "id": self.prospect.pk,
                "workstream": Prospect.Workstream.LINKEDIN_OUTREACH,
            },
            company_response.json()["company"]["prospects"],
        )


class ProspectSectionApiTests(CompanyApiTestMixin, TestCase):
    def setUp(self):
        self.linkedin_user = self.make_user(
            "linkedin-sections@example.com",
            Profile.Role.LINKEDIN_OUTREACH,
            "Leena LinkedIn",
        )
        self.other_linkedin_user = self.make_user(
            "other-linkedin-sections@example.com",
            Profile.Role.LINKEDIN_OUTREACH,
            "Omar LinkedIn",
        )
        self.intern = self.make_user(
            "intern-sections@example.com",
            Profile.Role.INTERN,
            "Isha Intern",
        )
        self.manager = self.make_user(
            "manager-sections@example.com",
            Profile.Role.MANAGER,
            "Meera Manager",
        )
        self.prospect = self.make_prospect(
            self.linkedin_user,
            "Section API Advisory",
            "https://section-api.example.com",
            contact_email="",
            contact_linkedin_url="https://www.linkedin.com/in/section-api-contact",
            founder_account=Prospect.FounderAccount.KANDARP_SONI,
            linkedin_connection_status=Prospect.LinkedInConnectionStatus.NOT_SENT,
            personalization_note="Chicago GCC advisory practice.",
        )

    def patch_section(self, route_name, payload, *, user=None, prospect=None):
        user = user or self.linkedin_user
        prospect = prospect or self.prospect
        return self.client.generic(
            "PATCH",
            reverse(route_name, args=[prospect.pk]),
            data=json.dumps(payload),
            content_type="application/json",
            HTTP_AUTHORIZATION=self.bearer(user),
        )

    def test_prospect_sent_api_marks_true_and_audits(self):
        response = self.patch_section(
            "api_prospect_sent",
            {"prospect_sent": True},
        )

        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(response.json()["message"], "Prospect sent status updated.")
        self.assertTrue(response.json()["prospect"]["prospect_sent"])
        self.assertNotIn("interest_signal", response.json()["prospect"])
        self.prospect.refresh_from_db()
        self.assertTrue(self.prospect.prospect_sent)
        audit = ProspectUpdateAudit.objects.get(prospect=self.prospect)
        self.assertEqual(audit.previous_values, {"prospect_sent": False})
        self.assertEqual(audit.changed_values, {"prospect_sent": True})

    def test_prospect_sent_api_rejects_wrong_types_and_unrelated_fields(self):
        wrong_type = self.patch_section(
            "api_prospect_sent",
            {"prospect_sent": "true"},
        )
        unrelated = self.patch_section(
            "api_prospect_sent",
            {"interest_signal": "Interested"},
        )

        self.assertEqual(wrong_type.status_code, 400)
        self.assertIn("prospect_sent", wrong_type.json()["field_errors"])
        self.assertEqual(unrelated.status_code, 400)
        self.assertIn("Unsupported field(s): interest_signal", unrelated.json()["error"])
        self.assertEqual(ProspectUpdateAudit.objects.count(), 0)

    def test_scoped_api_requires_owner_or_manager(self):
        forbidden = self.patch_section(
            "api_interest_handoff",
            {"interest_signal": "Interested"},
            user=self.other_linkedin_user,
        )
        unauthenticated = self.client.generic(
            "PATCH",
            reverse("api_interest_handoff", args=[self.prospect.pk]),
            data=json.dumps({"interest_signal": "Interested"}),
            content_type="application/json",
        )
        manager_update = self.patch_section(
            "api_interest_handoff",
            {"interest_signal": "Interested"},
            user=self.manager,
        )

        self.assertEqual(forbidden.status_code, 403)
        self.assertEqual(unauthenticated.status_code, 401)
        self.assertEqual(manager_update.status_code, 200)

    def test_founder_escalation_requires_accepted_invitation_and_response(self):
        response = self.patch_section(
            "api_founder_linkedin",
            {"founder_escalation_required": True},
        )

        self.assertEqual(response.status_code, 400)
        field_error = response.json()["field_errors"]["founder_escalation_required"]
        self.assertIn("invitation as accepted", field_error)
        self.assertIn("prospect response or interest signal", field_error)
        self.assertEqual(ProspectUpdateAudit.objects.count(), 0)

    def test_founder_linkedin_api_escalates_after_interest_and_acceptance(self):
        interest_response = self.patch_section(
            "api_interest_handoff",
            {
                "material_shared": Prospect.MaterialShared.ONE_PAGE,
                "interest_signal": "Asked for a founder discussion.",
                "questions_for_founders": "Can the engagement start with a pilot?",
            },
        )
        founder_response = self.patch_section(
            "api_founder_linkedin",
            {
                "linkedin_connection_status": Prospect.LinkedInConnectionStatus.ACCEPTED,
                "founder_escalation_required": True,
                "founder_escalation_notes": "Founder should answer the pilot question.",
            },
        )

        self.assertEqual(interest_response.status_code, 200, interest_response.content)
        self.assertEqual(founder_response.status_code, 200, founder_response.content)
        self.prospect.refresh_from_db()
        self.assertEqual(
            self.prospect.linkedin_connection_status,
            Prospect.LinkedInConnectionStatus.ACCEPTED,
        )
        self.assertTrue(self.prospect.founder_escalation_required)
        self.assertEqual(ProspectUpdateAudit.objects.filter(prospect=self.prospect).count(), 2)

    def test_founder_linkedin_api_rejects_non_linkedin_prospect(self):
        intern_prospect = self.make_prospect(
            self.intern,
            "Intern Section Advisory",
            "https://intern-section.example.com",
        )
        response = self.patch_section(
            "api_founder_linkedin",
            {"founder_account": Prospect.FounderAccount.KANDARP_SONI},
            user=self.intern,
            prospect=intern_prospect,
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn("founder_account", response.json()["field_errors"])

    def test_interest_handoff_api_updates_all_fields_and_validates_dropdown(self):
        updated = self.patch_section(
            "api_interest_handoff",
            {
                "material_shared": Prospect.MaterialShared.ONE_PAGE_AND_DECK,
                "interest_signal": "Requested a delivery-model discussion.",
                "questions_for_founders": "Can FlexGCC support a five-person pilot?",
            },
        )
        invalid = self.patch_section(
            "api_interest_handoff",
            {"material_shared": "full_proposal"},
        )

        self.assertEqual(updated.status_code, 200, updated.content)
        self.assertEqual(
            updated.json()["prospect"]["material_shared"],
            Prospect.MaterialShared.ONE_PAGE_AND_DECK,
        )
        self.assertEqual(invalid.status_code, 400)
        self.assertIn("material_shared", invalid.json()["field_errors"])
        self.prospect.refresh_from_db()
        self.assertEqual(
            self.prospect.questions_for_founders,
            "Can FlexGCC support a five-person pilot?",
        )

    def test_follow_up_api_updates_status_action_date_comments_and_stage(self):
        response = self.patch_section(
            "api_follow_up",
            {
                "status": Prospect.Status.MEETING_TO_SCHEDULE,
                "next_action": "Send three meeting slots",
                "next_action_date": "2026-09-15",
                "comments": "Prospect prefers morning US Central time.",
            },
        )

        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(response.json()["prospect"]["next_action_date"], "2026-09-15")
        self.prospect.refresh_from_db()
        self.assertEqual(self.prospect.status, Prospect.Status.MEETING_TO_SCHEDULE)
        self.assertEqual(self.prospect.stage, Prospect.Stage.INTERESTED)
        self.assertEqual(self.prospect.next_action_date, date(2026, 9, 15))
        audit = ProspectUpdateAudit.objects.get(prospect=self.prospect)
        self.assertEqual(audit.changed_values["next_action_date"], "2026-09-15")

    def test_follow_up_api_enforces_action_pair_and_date_types(self):
        missing_date = self.patch_section(
            "api_follow_up",
            {"next_action": "Send meeting slots"},
        )
        invalid_date = self.patch_section(
            "api_follow_up",
            {"next_action_date": "15/09/2026"},
        )

        self.assertEqual(missing_date.status_code, 400)
        self.assertIn("next_action_date", missing_date.json()["field_errors"])
        self.assertEqual(invalid_date.status_code, 400)
        self.assertIn("next_action_date", invalid_date.json()["field_errors"])

    def test_follow_up_api_requires_complete_timezone_aware_meeting(self):
        incomplete = self.patch_section(
            "api_follow_up",
            {"status": Prospect.Status.MEETING_SCHEDULED},
        )
        naive_datetime = self.patch_section(
            "api_follow_up",
            {"meeting_scheduled_at": "2026-09-18T10:30:00"},
        )
        complete = self.patch_section(
            "api_follow_up",
            {
                "status": Prospect.Status.MEETING_SCHEDULED,
                "meeting_scheduled_at": "2026-09-18T10:30:00-04:00",
                "meeting_timezone": "America/Chicago",
                "meeting_participants": "Asha Rao, Kandarp Soni",
            },
        )

        self.assertEqual(incomplete.status_code, 400)
        self.assertIn("meeting_scheduled_at", incomplete.json()["field_errors"])
        self.assertIn("meeting_timezone", incomplete.json()["field_errors"])
        self.assertIn("meeting_participants", incomplete.json()["field_errors"])
        self.assertEqual(naive_datetime.status_code, 400)
        self.assertIn("UTC offset", naive_datetime.json()["field_errors"]["meeting_scheduled_at"])
        self.assertEqual(complete.status_code, 200, complete.content)
        self.prospect.refresh_from_db()
        self.assertEqual(self.prospect.stage, Prospect.Stage.FOUNDER_MEETING_BOOKED)
        self.assertTrue(timezone.is_aware(self.prospect.meeting_scheduled_at))

    def test_scoped_get_returns_only_section_fields(self):
        response = self.client.get(
            reverse("api_prospect_sent", args=[self.prospect.pk]),
            HTTP_AUTHORIZATION=self.bearer(self.linkedin_user),
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            set(response.json()["prospect"]),
            {
                "id",
                "company_id",
                "company_name",
                "workstream",
                "owner",
                "prospect_sent",
            },
        )

    def test_general_prospect_api_remains_backward_compatible_with_new_fields(self):
        response = self.client.generic(
            "PATCH",
            reverse("api_prospect_detail", args=[self.prospect.pk]),
            data=json.dumps(
                {
                    "material_shared": Prospect.MaterialShared.SHORT_DECK,
                    "interest_signal": "Asked for more information.",
                    "comments": "Updated through the general endpoint.",
                }
            ),
            content_type="application/json",
            HTTP_AUTHORIZATION=self.bearer(self.linkedin_user),
        )

        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(
            response.json()["prospect"]["material_shared"],
            Prospect.MaterialShared.SHORT_DECK,
        )
        self.assertEqual(
            response.json()["prospect"]["comments"],
            "Updated through the general endpoint.",
        )


class OutreachHistoryApiTests(CompanyApiTestMixin, TestCase):
    def setUp(self):
        self.owner = self.make_user(
            "history-owner@example.com",
            Profile.Role.INTERN,
            "Isha Intern",
        )
        self.other = self.make_user(
            "history-other@example.com",
            Profile.Role.INTERN,
            "Omar Intern",
        )
        self.manager = self.make_user(
            "history-manager@example.com",
            Profile.Role.MANAGER,
            "Meera Manager",
        )
        self.prospect = self.make_prospect(
            self.owner,
            "History Advisory",
            "https://history-advisory.example.com",
        )

    def history_request(self, method, payload=None, *, user=None, prospect=None):
        user = user or self.owner
        prospect = prospect or self.prospect
        return self.client.generic(
            method,
            reverse("api_prospect_outreach_list", args=[prospect.pk]),
            data=json.dumps(payload) if payload is not None else "",
            content_type="application/json",
            HTTP_AUTHORIZATION=self.bearer(user),
        )

    def detail_request(self, method, outreach, payload=None, *, user=None):
        user = user or self.owner
        return self.client.generic(
            method,
            reverse("api_outreach_detail", args=[outreach.pk]),
            data=json.dumps(payload) if payload is not None else "",
            content_type="application/json",
            HTTP_AUTHORIZATION=self.bearer(user),
        )

    def test_owner_creates_and_lists_history_without_nullable_owner_lock_join(self):
        payload = {
            "activity_type": Outreach.ActivityType.INITIAL_OUTREACH,
            "medium": Outreach.Medium.EMAIL,
            "outreach_date": timezone.localdate().isoformat(),
            "response": "Asked for a one-page overview.",
        }
        with CaptureQueriesContext(connection) as queries:
            created = self.history_request("POST", payload)

        self.assertEqual(created.status_code, 201, created.content)
        self.assertEqual(created.json()["outreach"]["sequence_number"], 1)
        self.assertEqual(created.json()["outreach"]["recorded_by"]["email"], self.owner.email)
        self.prospect.refresh_from_db()
        self.assertEqual(self.prospect.stage, Prospect.Stage.RESPONDED)

        prospect_selects = [
            query["sql"]
            for query in queries.captured_queries
            if 'FROM "outreach_prospect"' in query["sql"]
            and query["sql"].lstrip().upper().startswith("SELECT")
        ]
        self.assertTrue(prospect_selects)
        for prospect_select in prospect_selects:
            self.assertNotIn('JOIN "auth_user"', prospect_select)

        listed = self.history_request("GET")
        self.assertEqual(listed.status_code, 200)
        self.assertEqual(listed.json()["count"], 1)
        self.assertEqual(listed.json()["maximum"], 5)
        self.assertEqual(
            listed.json()["outreaches"][0]["activity_type_display"],
            "Initial outreach",
        )

    def test_create_requires_exact_fields_and_enforces_contact_and_date_rules(self):
        missing = self.history_request(
            "POST",
            {"response": "Incomplete request."},
        )
        phone_without_number = self.history_request(
            "POST",
            {
                "activity_type": Outreach.ActivityType.PHONE_CALL,
                "medium": Outreach.Medium.PHONE,
                "outreach_date": timezone.localdate().isoformat(),
            },
        )
        future = self.history_request(
            "POST",
            {
                "activity_type": Outreach.ActivityType.INITIAL_OUTREACH,
                "medium": Outreach.Medium.EMAIL,
                "outreach_date": (timezone.localdate() + timedelta(days=1)).isoformat(),
            },
        )
        unsupported = self.history_request(
            "POST",
            {
                "activity_type": Outreach.ActivityType.INITIAL_OUTREACH,
                "medium": Outreach.Medium.EMAIL,
                "outreach_date": timezone.localdate().isoformat(),
                "sequence_number": 4,
            },
        )

        self.assertEqual(missing.status_code, 400)
        self.assertEqual(
            set(missing.json()["field_errors"]),
            {"activity_type", "medium", "outreach_date"},
        )
        self.assertEqual(phone_without_number.status_code, 400)
        self.assertIn("medium", phone_without_number.json()["field_errors"])
        self.assertEqual(future.status_code, 400)
        self.assertIn("outreach_date", future.json()["field_errors"])
        self.assertEqual(unsupported.status_code, 400)
        self.assertIn("Unsupported field(s): sequence_number", unsupported.json()["error"])
        self.assertEqual(self.prospect.outreaches.count(), 0)

    def test_create_enforces_research_state_permissions_and_five_record_limit(self):
        payload = {
            "activity_type": Outreach.ActivityType.INITIAL_OUTREACH,
            "medium": Outreach.Medium.EMAIL,
            "outreach_date": timezone.localdate().isoformat(),
        }
        forbidden = self.history_request("POST", payload, user=self.other)
        research_prospect = self.make_prospect(
            self.owner,
            "Research History Advisory",
            "https://research-history.example.com",
            stage=Prospect.Stage.RESEARCH,
        )
        research = self.history_request(
            "POST",
            payload,
            prospect=research_prospect,
        )
        for sequence_number in range(1, 6):
            Outreach.objects.create(
                prospect=self.prospect,
                sequence_number=sequence_number,
                activity_type=Outreach.ActivityType.FOLLOW_UP,
                medium=Outreach.Medium.EMAIL,
                outreach_date=timezone.localdate(),
                recorded_by=self.owner,
            )
        full = self.history_request("POST", payload)

        self.assertEqual(forbidden.status_code, 403)
        self.assertEqual(research.status_code, 400)
        self.assertIn("Complete the required", research.json()["error"])
        self.assertEqual(full.status_code, 409)
        self.assertEqual(self.prospect.outreaches.count(), 5)

    def test_linkedin_create_enforces_medium_and_connection_sequence(self):
        linkedin_user = self.make_user(
            "history-linkedin@example.com",
            Profile.Role.LINKEDIN_OUTREACH,
            "Leena LinkedIn",
        )
        linkedin_prospect = self.make_prospect(
            linkedin_user,
            "LinkedIn History Advisory",
            "https://linkedin-history.example.com",
            contact_email="",
            contact_linkedin_url="https://www.linkedin.com/in/history-contact",
            founder_account=Prospect.FounderAccount.KANDARP_SONI,
            linkedin_connection_status=Prospect.LinkedInConnectionStatus.NOT_SENT,
            personalization_note="Chicago advisory practice.",
        )
        premature = self.history_request(
            "POST",
            {
                "activity_type": Outreach.ActivityType.INITIAL_OUTREACH,
                "medium": Outreach.Medium.LINKEDIN,
                "outreach_date": timezone.localdate().isoformat(),
            },
            user=linkedin_user,
            prospect=linkedin_prospect,
        )
        wrong_medium = self.history_request(
            "POST",
            {
                "activity_type": Outreach.ActivityType.CONNECTION_REQUEST,
                "medium": Outreach.Medium.EMAIL,
                "outreach_date": timezone.localdate().isoformat(),
            },
            user=linkedin_user,
            prospect=linkedin_prospect,
        )
        valid = self.history_request(
            "POST",
            {
                "activity_type": Outreach.ActivityType.CONNECTION_REQUEST,
                "medium": Outreach.Medium.LINKEDIN,
                "outreach_date": timezone.localdate().isoformat(),
            },
            user=linkedin_user,
            prospect=linkedin_prospect,
        )

        self.assertEqual(premature.status_code, 400)
        self.assertIn("activity_type", premature.json()["field_errors"])
        self.assertEqual(wrong_medium.status_code, 400)
        self.assertIn("medium", wrong_medium.json()["field_errors"])
        self.assertEqual(valid.status_code, 201, valid.content)
        linkedin_prospect.refresh_from_db()
        self.assertEqual(
            linkedin_prospect.linkedin_connection_status,
            Prospect.LinkedInConnectionStatus.REQUEST_SENT,
        )

    def test_detail_get_and_patch_are_scoped_partial_and_audited(self):
        outreach = Outreach.objects.create(
            prospect=self.prospect,
            sequence_number=1,
            activity_type=Outreach.ActivityType.INITIAL_OUTREACH,
            medium=Outreach.Medium.EMAIL,
            outreach_date=timezone.localdate() - timedelta(days=2),
            response="Original note.",
            recorded_by=self.owner,
        )
        new_date = timezone.localdate() - timedelta(days=1)
        updated = self.detail_request(
            "PATCH",
            outreach,
            {
                "outreach_date": new_date.isoformat(),
                "response": "Corrected factual note.",
            },
        )
        forbidden = self.detail_request("GET", outreach, user=self.other)
        manager_read = self.detail_request("GET", outreach, user=self.manager)
        immutable = self.detail_request(
            "PATCH",
            outreach,
            {"sequence_number": 5},
        )

        self.assertEqual(updated.status_code, 200, updated.content)
        self.assertEqual(updated.json()["outreach"]["outreach_date"], new_date.isoformat())
        self.assertEqual(updated.json()["outreach"]["activity_type"], Outreach.ActivityType.INITIAL_OUTREACH)
        self.assertEqual(updated.json()["modified_by"]["email"], self.owner.email)
        self.assertEqual(forbidden.status_code, 403)
        self.assertEqual(manager_read.status_code, 200)
        self.assertEqual(immutable.status_code, 400)
        self.assertIn("Unsupported field(s): sequence_number", immutable.json()["error"])

        outreach.refresh_from_db()
        self.assertEqual(outreach.response, "Corrected factual note.")
        audit = OutreachUpdateAudit.objects.get(outreach=outreach)
        self.assertEqual(audit.modified_by, self.owner)
        self.assertEqual(audit.previous_values["response"], "Original note.")
        self.assertEqual(audit.changed_values["outreach_date"], new_date.isoformat())

    def test_history_apis_require_a_bearer_token(self):
        list_response = self.client.get(
            reverse("api_prospect_outreach_list", args=[self.prospect.pk])
        )
        missing_response = self.client.get(
            reverse("api_outreach_detail", args=[999999]),
            HTTP_AUTHORIZATION=self.bearer(self.manager),
        )

        self.assertEqual(list_response.status_code, 401)
        self.assertEqual(missing_response.status_code, 404)


class ApiDocumentationTests(CompanyApiTestMixin, TestCase):
    def setUp(self):
        self.user = self.make_user(
            "api-docs@example.com",
            Profile.Role.MANAGER,
            "Meera Manager",
        )

    def test_api_access_lists_separate_guides(self):
        self.client.force_login(self.user)
        response = self.client.get(reverse("api_access"))

        self.assertEqual(response.status_code, 200)
        for slug, title in [
            ("prospect-sent", "Prospect sent API"),
            ("founder-linkedin", "Founder LinkedIn account API"),
            ("interest-handoff", "Interest and handoff API"),
            ("follow-up", "Follow-up API"),
            ("outreach-history-list-create", "Outreach history list and create API"),
            ("outreach-history-detail", "Outreach history detail API"),
        ]:
            with self.subTest(slug=slug):
                self.assertContains(response, title)
                self.assertContains(response, reverse("api_documentation", args=[slug]))

    def test_each_guide_documents_fields_types_values_and_examples(self):
        self.client.force_login(self.user)
        expectations = {
            "prospect-sent": ["prospect_sent", "boolean", '"prospect_sent": true'],
            "founder-linkedin": [
                "founder_escalation_required",
                "kandarp_soni",
                "sunit_kala",
                "accepted",
            ],
            "interest-handoff": [
                "material_shared",
                "one_page",
                "short_deck",
                "one_page_and_deck",
            ],
            "follow-up": [
                "next_action_date",
                "meeting_scheduled",
                "meeting_scheduled_at",
                "ISO 8601",
            ],
            "outreach-history-list-create": [
                "activity_type",
                "connection_request",
                "outreach_date",
                "GET, POST",
                "curl --request POST",
            ],
            "outreach-history-detail": [
                "outreach_id",
                "sequence_number",
                "previous values",
                "GET, PATCH",
            ],
        }

        for slug, expected_text in expectations.items():
            with self.subTest(slug=slug):
                response = self.client.get(reverse("api_documentation", args=[slug]))
                self.assertEqual(response.status_code, 200)
                self.assertIn("private", response["Cache-Control"])
                self.assertIn("no-store", response["Cache-Control"])
                for text in expected_text:
                    self.assertContains(response, escape(text))
                self.assertContains(response, "Authorization: Bearer")
                if slug != "outreach-history-list-create":
                    self.assertContains(response, "curl --request PATCH")

    def test_documentation_pages_require_login_and_unknown_page_is_404(self):
        protected = self.client.get(
            reverse("api_documentation", args=["prospect-sent"])
        )
        self.client.force_login(self.user)
        missing = self.client.get(
            reverse("api_documentation", args=["does-not-exist"])
        )

        self.assertEqual(protected.status_code, 302)
        self.assertEqual(missing.status_code, 404)
